# -*- coding: utf-8 -*-
"""poster_kit — 2026 원탐학회 가뭄 포스터 그림·표 한 벌 (S1 · 2026-10-15).

수치 기준: gee docs/DROUGHT_2026_EVIDENCE_KR.md 「포스터 조립용 최종 목록」(v11 · EGM 켬).
서사 순서: ① 이례성(WAMIS 댐별 10년 기록·정부 예·경보) → ② SAR 면적 감소 → ③ 공식 자료와 일치(하한) → ④ 부록.

판단 기준 (사용자 2026-10-15) — 자료마다 자기 기준으로 판단한다. 순위로 가뭄 권역을 정하지 않는다.
    WAMIS        댐마다 자기 10년(2017~2026) 8월 월평균 저수위와 비교 — 중앙값 대비(m), 10년 최저, 더 낮았던 해
    정부 보도자료  보도자료 자체의 판단(행안부 8월 가뭄 예·경보 단계, 기후부 진입일) — 그것도 평년·예년 기준
    SAR          판정이 아니라 관측(전년도 대비 수체 면적, 저수량 감소의 하한)

    그림1  WAMIS 8월 평균 저수위 10년 꺾은선 (8개 저수지)
    그림2  남부 지도 — (a) 댐별 10년 기록 + 예·경보 단계 (b) SAR 댐유역 8월 Δ% (같은 유역 폴리곤)
    그림3  8개 댐유역 지도·막대          (poster_sar_figs 그림A)
    그림4  SAR 면적 ↔ WAMIS 저수량·민감도 (poster_sar_figs 그림B)
    그림5  SAR·S2 화소 대응              (poster_sar_figs 그림C)
    그림6  관측쌍 타임라인 + 직전 5일 강우 (방법)
    그림7  부록 — 강우 위상: 저수지 ↔ 유역
    표1~11  PNG + 표모음_편집용.xlsx + 표모음_A0_편집용.pptx(포스터와 같은 A0 세로·22/20 pt)
    00_S1_SAR_포스터_그림표_정리.md — 그림·표·캡션·주의를 한 문서로 (표는 PNG 와 같은 값에서 생성)

입력은 전부 읽기만 한다. 유역 정본 배정은 gee scripts/canon_table.py 의 규칙을 그대로 옮겨
(그 스크립트는 import 시 파일을 쓰므로 부르지 않는다) output/canon_table.txt 와 대조한다.
WAMIS 월자료(남부 33개 댐, 조정지 제외)는 --out/_자료 에 받아 둔다(없으면 받는다 — 인증키 없는 8080 HTTP).

실행 (gis_clean; xlsx·pptx 는 openpyxl·python-pptx 가 있는 경로를 --pylib 로)
    $env:GDAL_DATA = "F:\\envs\\gis_clean\\Library\\share\\gdal"
    F:\\envs\\gis_clean\\python.exe scripts\\poster_kit.py --out <폴더> [--pylib <site-packages>]
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import matplotlib
import numpy as np
import pandas as pd

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import matplotlib.ticker as mticker  # noqa: E402
from matplotlib.colors import BoundaryNorm, ListedColormap  # noqa: E402
from matplotlib.lines import Line2D  # noqa: E402
from matplotlib.patches import Patch, Rectangle  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parent))
import poster_sar_figs as psf  # noqa: E402

ROOT, GEE, WAMIS_DIR = psf.ROOT, psf.GEE, psf.WAMIS
S1 = ROOT / "S1"
RANK_CSV = S1 / "data" / "wamis_rank_jul_aug.csv"          # 남부 댐 목록·코드·관리 기관만 읽는다(순위 열은 안 쓴다)
RAIN_CSV = S1 / "data" / "rain_phase_5d.csv"
POINTS_JSON = S1 / "reservoir_points.json"
S2_DIR = Path(r"E:\~K-water_지상국\20261021_원탐학회\S2_자료_20261015")
MON_NAME = "wamis_monthly_south_2017_2026.csv"
FS = psf.FS
C_JUL, C_AUG, C_NAVY, C_INK, C_MUTED, C_GRID = psf.C_JUL, psf.C_AUG, psf.C_NAVY, psf.C_INK, psf.C_MUTED, psf.C_GRID
C_25 = "#3a7dbf"                     # 2025 (전년도) — 파랑
C_26 = C_AUG                         # 2026 — 진주황
C_HI = "#c4400a"                     # 표 강조 숫자
C_UP = "#2171b5"                     # 증가·중앙값 위
M = "−"


def f1(v: float, sign: bool = True) -> str:
    """소수 1자리, U+2212."""
    s = f"{v:+.1f}" if sign else f"{v:.1f}"
    return s.replace("-", M) if v < 0 else (s if v > 0 or not sign else s.replace("+", ""))


def f2(v: float, sign: bool = False) -> str:
    s = f"{v:+.2f}" if sign else f"{v:.2f}"
    return s.replace("-", M)


# ════════════════════════ SAR 정본 배정 ════════════════════════
# gee scripts/canon_table.py 2026-10-15 판 규칙 (RUNS·ASSIGN·_fix 우선·동급 겹침 정지)
RUNS = {
    "7월A": ("0725↔0715", "ASC54", ["watershed_change_v11_jul_a.csv", "watershed_change_v11_jul_a_han.csv",
                                    "watershed_change_v11_jul_a_fix.csv", "watershed_change_v11_jul_a_namgang.csv"]),
    "7월B": ("0718↔0720", "DESC134", ["watershed_change_v11_jul_b.csv"]),
    "8월A": ("0806↔0802", "ASC54", ["watershed_change_v11_aug_a.csv", "watershed_change_v11_aug_a_fix.csv"]),
    "8월B": ("0811↔0807", "ASC127", ["watershed_change_v11_aug_b.csv", "watershed_change_v11_aug_b_fix.csv"]),
}
ASSIGN = {("7월", "대청"): "7월B", ("7월", "남강"): "7월B", ("7월", "영주"): "7월A", ("7월", "광동"): "7월A",
          ("7월", "군위"): "7월A", ("7월", "부항"): "7월A", ("7월", "합천"): "7월A",
          ("8월", "대청"): "8월B", ("8월", "합천"): "8월A", ("8월", "부항"): "8월A", ("8월", "남강"): "8월B"}
HAN = {"한 강", "한강동해", "한강서해", "충주", "횡성", "광동", "달방", "소양강", "임남댐", "평화화천", "한탄강", "군남"}
OUT_OF_SCOPE = HAN | {"제주도", "삽교천", "안성천"}     # 포스터 = 남부 15 대권역 (사용자 10-14)


def _load_run(run: str) -> dict:
    out, grade = {}, {}
    for fn in sorted(RUNS[run][2], key=lambda f: 2 if "_fix" in f else 1):
        p = GEE / fn
        if not p.exists():
            continue
        g = 2 if "_fix" in fn else 1
        for r in pd.read_csv(p, encoding="utf-8-sig", dtype=str).to_dict("records"):
            k = (r["구분"].strip(), r["대상"].strip())
            if k in out and grade[k] == g:
                raise SystemExit(f"동급 겹침 {k} — {fn}")
            out[k], grade[k] = r, g
    return out


def canon() -> pd.DataFrame:
    data = {r: _load_run(r) for r in RUNS}
    recs = []
    for kind, nm in sorted({k for r in data for k in data[r]}):
        if nm in OUT_OF_SCOPE:
            continue
        for month in ("7월", "8월"):
            c = [r for r in RUNS if r.startswith(month) and (kind, nm) in data[r]]
            if not c:
                continue
            run = c[0] if len(c) == 1 else ASSIGN[(month, nm)]
            r = data[run][(kind, nm)]
            recs.append(dict(구분=kind, 대상=nm, 월=month, 회차=run, 쌍=RUNS[run][0], 궤도=RUNS[run][1],
                             pct=float(r["Δ%"]), w25=float(r["2025수체km2"]), w26=float(r["2026수체km2"]),
                             lost=float(r["소실km2"]), new=float(r["신규km2"]), cov=float(r["커버"])))
    df = pd.DataFrame(recs)
    _check_canon(df)
    return df


def _check_canon(df: pd.DataFrame) -> None:
    """gee output/canon_table.txt(사람이 읽은 정본) 와 Δ% 를 칸마다 대조한다."""
    p = GEE / "canon_table.txt"
    if not p.exists():
        print("⚠ canon_table.txt 없음 — 대조 생략"); return
    n = 0
    for line in p.read_text(encoding="utf-8").splitlines():
        if not line.startswith(("대권역", "댐유역")) or "|" not in line:
            continue
        left, *cells = line.split("|")
        kind, nm = left.split(None, 1)[0], left.split(None, 1)[1].strip()
        if nm in OUT_OF_SCOPE:
            continue
        for month, cell in zip(("7월", "8월"), cells):
            m = re.match(r"\s*(\d월[AB])\s+([+-]?\d+\.\d)", cell)
            if not m:
                continue
            row = df[(df.구분 == kind) & (df.대상 == nm) & (df.월 == month)].iloc[0]
            assert row.회차 == m.group(1) and abs(row.pct - float(m.group(2))) < 0.05, (kind, nm, month)
            n += 1
    print(f"정본 대조: canon_table.txt 와 {n}칸 일치")


def cget(cn: pd.DataFrame, nm: str, month: str, kind: str = "댐유역") -> pd.Series | None:
    s = cn[(cn.구분 == kind) & (cn.대상 == nm) & (cn.월 == month)]
    return None if s.empty else s.iloc[0]


# 대권역 — (표시 이름, 대권역 shp BBSNNM, WAMIS 권역 이름)
REG = [("낙동강", "낙동강", "낙동강"), ("금강", "금 강", "금강"), ("섬진강", "섬진강", "섬진강"),
       ("영산강", "영산강", "영산강"), ("태화강", "태화강", "태화강"), ("낙동강남해", "낙동강남해", "낙동강남해"),
       ("섬진강남해", "섬진강남해", "섬진강남해"), ("만경·동진", "만경동진", "만경 · 동진"),
       ("낙동강동해", "낙동강동해", "낙동강동해"), ("탐진강", "탐진강", "탐진강"),
       ("회야·수영", "회야수영", "회야 · 수영"), ("형산강", "형산강", "형산강")]
NO_DAM = ["영산강남해", "금강서해", "영산강서해"]
WREG_ORDER = ["낙동강", "형산강", "태화강", "회야 · 수영", "낙동강동해", "낙동강남해", "금강", "금강서해", "만경 · 동진",
              "섬진강", "섬진강남해", "영산강", "탐진강"]


# ════════════════════════ WAMIS — 댐별 10년 기록 ════════════════════════
WAMIS_NAME = {"주암": "주암(본)", "부항": "김천부항"}
NO_WAMIS_REG = {"군위": "낙동강", "보현산": "낙동강", "보령": "금강서해"}     # WAMIS 월자료 없는 SAR 댐유역
RES_DAM = {"안동호": "안동", "임하호": "임하", "밀양호": "밀양", "영천호": "영천", "성덕저수지": "성덕",
           "운문호": "운문", "옥정호": "섬진강", "평림호": "평림"}
MGR = {"한국수자원공사": "K-water", "한국수력원자력": "한수원", "한국농어촌공사": "농어촌공사"}


def wamis_south(cache: Path) -> pd.DataFrame:
    """남부 WAMIS 댐(조정지 제외) 7·8월 월평균 저수위 2017~2026. 없으면 받는다."""
    if not cache.exists():
        sys.path.insert(0, str(S1))
        from s1.tools.wamis.wamis_august_climatology import monthly
        rk = pd.read_csv(RANK_CSV, encoding="utf-8-sig", dtype={"코드": str})
        rk["권역"] = rk["권역"].fillna("")
        south = rk[~rk["권역"].isin(["한강", "한강동해", ""]) & ~rk["댐"].str.contains("조정지")]
        rows = []
        for _, r in south.iterrows():
            for (y, mm), v in sorted(monthly(r["코드"], ("07", "08"), spans=((2009, 2016), (2017, 2026))).items()):
                rows.append({"권역": r["권역"], "댐": r["댐"], "코드": r["코드"],
                             "관리": r["관리"] if isinstance(r["관리"], str) else "", "연": y, "월": int(mm),
                             "월평균수위_m": v})
        cache.parent.mkdir(parents=True, exist_ok=True)
        pd.DataFrame(rows).to_csv(cache, index=False, encoding="utf-8-sig")
    df = pd.read_csv(cache, encoding="utf-8-sig", dtype={"코드": str})
    df["관리"] = df["관리"].fillna("")
    return df


def wstat(mon: pd.DataFrame, dam: str, month: int = 8) -> dict | None:
    """댐 하나를 자기 10년 기록과 비교한다 — 다른 댐과 견주지 않는다."""
    g = mon[(mon.댐 == WAMIS_NAME.get(dam, dam)) & (mon.월 == month) & (mon.연 >= 2017)]
    if g.empty:
        return None
    s = g.set_index("연")["월평균수위_m"].sort_index()
    if 2026 not in s.index:
        return None
    med = float(s.median())
    return dict(s=s, val=float(s[2026]), med=med, anom=float(s[2026] - med),
                a25=float(s[2025] - med) if 2025 in s.index else np.nan, n=len(s),
                lower=[int(y) for y in s.index if s[y] < s[2026] - 1e-9],
                region=g["권역"].iloc[0], mgr=MGR.get(g["관리"].iloc[0], "—"))


def where(st: dict | None) -> str:
    if st is None:
        return "월자료 없음"
    if not st["lower"]:
        return f"{st['n']}년 최저"
    if abs(st["anom"]) < 0.005:
        return "중앙값 수준"
    if st["anom"] > 0:
        return "중앙값 위"
    if len(st["lower"]) <= 2:
        return "·".join(str(y) for y in st["lower"]) + "년 다음으로 낮음"
    return "중앙값 아래"


def anom_s(st: dict | None, key: str = "anom") -> str:
    return "—" if st is None or not np.isfinite(st[key]) else f2(st[key], True)


# 본문 8개 — 순서·이름·단계
DAMS8 = psf.DAMS
RES8 = psf.RESV
REGION8 = {"안동": "낙동강", "임하": "낙동강", "밀양": "낙동강", "영천": "낙동강", "성덕": "낙동강",
           "운문": "낙동강", "섬진강": "섬진강", "평림": "영산강"}
STAGE = psf.STAGE                       # 행안부 「8월 가뭄 예·경보」(2026-08-10 발표) — 원문 대조 완료
STAGE_ENTRY = {"운문": "07-15", "안동": "07-08", "임하": "07-08", "성덕": "07-18", "밀양": "08-02",
               "평림": "08-02"}         # 기후부 보도참고자료 원문 날짜
SMALL_DKM2 = psf.SMALL_DKM2             # 성덕 −0.25 · 평림 −0.03 (Δ% 에서 되돌린 값)


# ════════════════════════ 그림 ════════════════════════
def fig1(mon: pd.DataFrame, out: Path) -> None:
    """8개 저수지 8월 평균 저수위 10년 꺾은선 — 칸마다 자기 축."""
    fig, axs = plt.subplots(2, 4, figsize=(14.8, 7.6))
    fig.subplots_adjust(left=0.055, right=0.99, top=0.84, bottom=0.07, wspace=0.38, hspace=0.50)
    C_LINE = "#8a9bb0"
    for ax, res in zip(axs.ravel(), RES_DAM):
        st = wstat(mon, RES_DAM[res])
        s = st["s"]
        lo, hi = s.min(), s.max()
        span = hi - lo
        ax.axhline(st["med"], color="#9a9a9a", lw=1.0, ls=(0, (3, 3)), zorder=1)
        ax.plot(s.index, s.values, color=C_LINE, lw=2.0, marker="o", ms=5, mfc="white", mec=C_LINE, mew=1.4, zorder=2)
        ax.plot([2025, 2026], [s[2025], s[2026]], color=C_26, lw=2.8, zorder=3)
        ax.scatter([2025], [s[2025]], s=85, color=C_25, edgecolor="white", lw=1.2, zorder=4)
        ax.scatter([2026], [s[2026]], s=120, marker="s", color=C_26, edgecolor="white", lw=1.2, zorder=5)
        lw_ = where(st)
        lab = lw_ if lw_.endswith("최저") or lw_.endswith("낮음") else f"중앙값보다 {abs(st['anom']):.1f} m 낮음"
        ax.text(2026.15, lo - 0.07 * span, f"2026  {s[2026]:.1f} m\n{lab}", ha="right", va="top", fontsize=FS - 5,
                color=C_26, fontweight="bold" if lw_.endswith("최저") else "normal", linespacing=1.15, zorder=6)
        ax.set_ylim(lo - 0.42 * span, hi + 0.12 * span)
        ax.set_xlim(2016.5, 2026.6)
        ax.set_xticks([2017, 2020, 2023, 2026]); ax.set_xticklabels(["2017", "2020", "2023", "2026"], fontsize=FS - 5)
        ax.tick_params(axis="y", labelsize=FS - 5)
        ax.yaxis.set_major_locator(mticker.MaxNLocator(4, integer=True))
        ax.grid(axis="y", color=C_GRID, lw=0.7); ax.set_axisbelow(True)
        for s_ in ("top", "right"):
            ax.spines[s_].set_visible(False)
        ax.set_title(res, loc="left", fontsize=FS - 2, fontweight="bold", color=C_INK, pad=6)
        ax.set_title(f"전년 대비 {f1(s[2026] - s[2025])} m", loc="right", fontsize=FS - 5, color=C_MUTED, pad=7)
    for ax in axs[:, 0]:
        ax.set_ylabel("EL.m", fontsize=FS - 4, color=C_MUTED)
    fig.legend(handles=[Line2D([], [], color=C_LINE, lw=2, marker="o", mfc="white", mec=C_LINE),
                        Line2D([], [], marker="o", ls="", ms=10, color=C_25),
                        Line2D([], [], marker="s", ls="", ms=10, color=C_26),
                        Line2D([], [], color="#9a9a9a", lw=1.0, ls=(0, (3, 3)))],
               labels=["8월 평균 저수위 (2017~2026)", "2025 (전년도)", "2026", "10년 중앙값"],
               loc="upper left", bbox_to_anchor=(0.045, 0.985), ncol=4, fontsize=FS - 2, handlelength=1.8,
               columnspacing=1.6)
    fig.text(0.99, 0.985, "WAMIS 월자료 · 칸마다 세로축이 다르다 · 옥정호 = 섬진강댐", ha="right", va="top",
             fontsize=FS - 5, color=C_MUTED)
    fig.savefig(out, facecolor="white"); plt.close(fig)


DCOL = ListedColormap(["#7f2704", "#a63603", "#d94801", "#f16913", "#fd8d3c", "#fdd0a2",
                       "#c6dbef", "#6baed6", "#2171b5"])
DBND = [-50, -25, -20, -15, -10, -5, 0, 5, 10, 20]
WCAT = ["#a63603", "#fd8d3c", "#9ecae1", "#e6e6e6"]       # 자기 기록 최저 · 중앙값 아래 · 위/수준 · 자료 없음


def _wcat(st: dict | None) -> int:
    if st is None:
        return 3
    if not st["lower"]:
        return 0
    return 1 if st["anom"] < -0.005 else 2


def _spread(items: list, ymin: float, ymax: float, dy: float) -> list:
    """(이름, cx, cy) 를 cy 내림차순으로 놓고 겹치지 않게 라벨 y 를 아래로 민다."""
    items = sorted(items, key=lambda t: -t[2])
    ys, last = [], ymax + dy
    for _, _, cy in items:
        y = min(cy, last - dy)
        ys.append(y); last = y
    shift = max(0.0, ymin - ys[-1]) if ys else 0.0
    return [(it[0], it[1], it[2], y + shift) for it, y in zip(items, ys)]


def _labels(ax, items, x0, x1, y0, y1, text_of, style_of, dy):
    east = [t for t in items if t[1] > 1_060_000]
    west = [t for t in items if t[1] <= 1_060_000]
    for grp, xl, ha in ((east, x1 + 18_000, "left"), (west, x0 - 18_000, "right")):
        for nm, cx, cy, ly in _spread(grp, y0 + 20_000, y1 - 10_000, dy):
            fs, col, wt = style_of(nm)
            ax.annotate(text_of(nm), xy=(cx, cy), xytext=(xl, ly), ha=ha, va="center", fontsize=fs, color=col,
                        fontweight=wt, arrowprops=dict(arrowstyle="-", color="#777777", lw=0.7, shrinkA=1, shrinkB=0))


def fig2(cn: pd.DataFrame, mon: pd.DataFrame, out: Path) -> None:
    import geopandas as gpd
    bb = gpd.read_file(psf.SHP_BBSN).to_crs(5179)
    dm = gpd.read_file(psf.SHP_DAM).to_crs(5179)
    south_nm = [r[1] for r in REG] + NO_DAM
    south = bb[bb["BBSNNM"].isin(south_nm)]
    other = bb[~bb["BBSNNM"].isin(south_nm) & (bb["BBSNNM"] != "제주도")]
    big4 = bb[bb["BBSNCD"].astype(int).isin([20, 30, 40, 50])]
    aug = cn[(cn.구분 == "댐유역") & (cn.월 == "8월")].set_index("대상")
    dm = dm[dm["damname"].isin(aug.index)].copy()
    dm["pct"] = dm["damname"].map(aug["pct"])
    dm["cat"] = dm["damname"].map(lambda d: _wcat(wstat(mon, d)))
    dm = dm.assign(area=dm.area).sort_values("area", ascending=False)
    parts = south.explode(index_parts=False)
    x0, y0, x1, y1 = parts[parts.geometry.centroid.x < 1_250_000].total_bounds   # 울릉·독도 제외한 범위
    # 유역 폴리곤이 없는 WAMIS 댐 — 담수호 좌표로 점을 찍는다
    pts = json.loads(POINTS_JSON.read_text(encoding="utf-8"))["lakes"]
    extra = {"동복": "동복호", "담양": "담양호", "장성": "장성호", "나주": "나주호"}
    pdf = gpd.GeoDataFrame({"dam": list(extra)},
                           geometry=gpd.points_from_xy([pts[v]["lon"] for v in extra.values()],
                                                       [pts[v]["lat"] for v in extra.values()]), crs=4326).to_crs(5179)
    pdf["cat"] = pdf["dam"].map(lambda d: _wcat(wstat(mon, d)))

    # (a) 는 서쪽 라벨(대청·용담 「단계 없음」)이 길어 서쪽 여백을 더 준다 — 칸 폭을 x 범위에 비례시켜 축척을 맞춘다
    xr = [(x0 - 235_000, x1 + 165_000), (x0 - 235_000, x1 + 150_000)]
    fig = plt.figure(figsize=(14.8, 6.8))
    gs = fig.add_gridspec(1, 2, wspace=0.0, left=0.0, right=1.0, top=0.86, bottom=0.15,
                          width_ratios=[b - a for a, b in xr])
    axs = [fig.add_subplot(gs[0, i]) for i in range(2)]
    for ax, (xa, xb) in zip(axs, xr):
        other.plot(ax=ax, color="#f3f3f3", edgecolor="#d6d6d6", lw=0.5)
        south.plot(ax=ax, color="#eef1f5", edgecolor="#b9c3d1", lw=0.6)
        big4.boundary.plot(ax=ax, color=C_NAVY, lw=1.1)
        ax.set_xlim(xa, xb); ax.set_ylim(y0 - 10_000, y1 + 15_000)
        ax.set_aspect("equal"); ax.set_axis_off()
    main = dm[dm["damname"].isin(DAMS8)]

    # (a) WAMIS 댐별 10년 기록 + 정부 단계
    ax = axs[0]
    has = dm[dm["cat"] < 3]
    has.plot(ax=ax, color=[WCAT[c] for c in has["cat"]], edgecolor="#555555", lw=0.5)
    dm[dm["cat"] == 3].plot(ax=ax, facecolor="#f2f2f2", edgecolor="#9a9a9a", lw=0.5, hatch="///")
    pdf.plot(ax=ax, color=[WCAT[c] for c in pdf["cat"]], edgecolor="#333333", markersize=70, lw=0.8, zorder=5)
    main.boundary.plot(ax=ax, color="#111111", lw=1.6)
    show = DAMS8 + ["용담", "대청", "영주", "안계"]
    items = [(r["damname"], *r.geometry.representative_point().coords[0])
             for _, r in dm[dm["damname"].isin(show)].iterrows()]

    def ta(nm):
        st = wstat(mon, nm)
        stg = STAGE.get(nm, "단계 없음" if nm in ("용담", "대청") else "")
        return f"{nm} {stg} {f1(st['anom'])}".replace("  ", " ")

    def sa(nm):
        return (FS - 3, C_INK, "bold") if nm in DAMS8 else (FS - 4, C_MUTED, "normal")

    _labels(ax, items, x0, x1, y0, y1, ta, sa, (y1 - y0) / 12)
    ax.set_title("(a) WAMIS 2026년 8월 저수위 — 댐마다 자기 10년 기록과 비교", loc="left", fontweight="bold",
                 fontsize=FS - 1, pad=24)
    ax.text(0.0, 1.01, "라벨 = 8월 가뭄 예·경보 단계(행안부 8.10) · 2026-08 저수위의 10년 중앙값 대비(m)", transform=ax.transAxes,
            fontsize=FS - 5, color=C_MUTED, va="bottom")
    ax.legend(handles=[Patch(color=WCAT[0]), Patch(color=WCAT[1]), Patch(color=WCAT[2]),
                       Patch(fc="#f2f2f2", hatch="///", ec="#9a9a9a")],
              labels=["자기 기록 최저", "10년 중앙값 아래", "중앙값 위·수준", "WAMIS 월자료 없음"],
              loc="upper left", bbox_to_anchor=(0.02, -0.005), ncol=2, fontsize=FS - 4, handlelength=1.2,
              columnspacing=1.2)

    # (b) SAR 댐유역 8월 Δ%
    ax = axs[1]
    dnorm = BoundaryNorm(DBND, DCOL.N)
    dm.plot(ax=ax, color=[DCOL(dnorm(v)) for v in dm["pct"]], edgecolor="#555555", lw=0.5)
    main.boundary.plot(ax=ax, color="#111111", lw=1.6)
    note = {"대청": " (변화 없음)", "평림": "†", "성덕": "†"}
    _labels(ax, items, x0, x1, y0, y1, lambda nm: f"{nm} {f1(aug.loc[nm, 'pct'])}{note.get(nm, '')}", sa,
            (y1 - y0) / 12)
    ax.set_title("(b) SAR 댐유역 수체 면적 Δ% — 2025→2026 8월", loc="left", fontweight="bold", fontsize=FS - 1,
                 pad=24)
    ax.text(0.0, 1.01, "같은 상대궤도 1:1 쌍 · 굵은 테 = 본문 8개 댐유역", transform=ax.transAxes,
            fontsize=FS - 5, color=C_MUTED, va="bottom")
    cax = fig.add_axes([0.56, 0.085, 0.40, 0.025])
    cb = fig.colorbar(plt.cm.ScalarMappable(norm=dnorm, cmap=DCOL), cax=cax, orientation="horizontal")
    cb.set_ticks(DBND); cb.set_ticklabels([str(b).replace("-", M) for b in DBND])
    cb.outline.set_visible(False); cb.ax.tick_params(labelsize=FS - 5, length=0)
    cb.set_label("Δ% (2025→2026, 8월)", fontsize=FS - 4, color=C_MUTED)
    fig.savefig(out, facecolor="white", bbox_inches="tight", pad_inches=0.12); plt.close(fig)


PAIRS = [("7월A", 54, "20250725", "20260715", "main"), ("7월B", 134, "20250718", "20260720", "main"),
         ("8월A", 54, "20250806", "20260802", "main"), ("8월B", 127, "20250811", "20260807", "main"),
         ("민감도", 134, "20250811", "20260813", "sens"), ("민감도", 134, "20250811", "20260825", "sens"),
         ("민감도", 134, "20250823", "20260813", "sens"), ("민감도", 134, "20250823", "20260825", "sens"),
         ("건조 대조", 54, "20250812", "20260802", "dry")]
LANES = [(54, "ASC54"), (134, "DESC134"), (127, "ASC127")]
SAT_MK = {"S1A": "^", "S1C": "o", "S1D": "s"}
BASINS5 = ["한강", "금강", "낙동강", "섬진강", "영산강"]
PAIR_BASINS = {"7월A": ["낙동강"], "8월A": ["낙동강"], "건조 대조": ["낙동강"], "7월B": ["섬진강", "영산강"],
               "8월B": ["섬진강", "영산강"], "민감도": ["낙동강", "섬진강", "영산강"]}


def fig6(rain: pd.DataFrame, out: Path) -> None:
    def X(d: str) -> int:
        return (date(2000, int(d[4:6]), int(d[6:8])) - date(2000, 7, 1)).days

    lane_y = {}
    for i, (rel, _) in enumerate(LANES):
        lane_y[(rel, "2025")] = -i * 2.7
        lane_y[(rel, "2026")] = -i * 2.7 - 1.0
    fig, ax = plt.subplots(figsize=(14.8, 6.4))
    fig.subplots_adjust(left=0.13, right=0.985, top=0.83, bottom=0.11)
    for (rel, yr), y in lane_y.items():
        ax.plot([0, 61], [y, y], color="#e3e3e3", lw=1.0, zorder=0)
    used = {(p[1], p[2]) for p in PAIRS} | {(p[1], p[3]) for p in PAIRS}
    style = {"main": dict(color=C_AUG, lw=3.0, ls="-", alpha=0.9),
             "sens": dict(color=C_NAVY, lw=1.3, ls=(0, (4, 2)), alpha=0.8),
             "dry": dict(color="#555555", lw=1.6, ls=(0, (1, 1.5)), alpha=0.9)}
    labeled = set()
    for lab, rel, d25, d26, kind in PAIRS:
        xa, ya, xb, yb = X(d25), lane_y[(rel, "2025")], X(d26), lane_y[(rel, "2026")]
        ax.plot([xa, xb], [ya, yb], zorder=1, **style[kind])
        if lab not in labeled:
            labeled.add(lab)
            tx, ty = (xa + xb) / 2, (ya + yb) / 2
            off = {"7월A": (-1.2, 0), "7월B": (-1.6, 0), "8월A": (0.9, 0), "8월B": (0.9, 0),
                   "민감도": (6.2, 0.05), "건조 대조": (3.4, -0.05)}[lab]
            ax.text(tx + off[0], ty + off[1], lab, ha="center", va="center", fontsize=FS - 3,
                    color=style[kind]["color"], fontweight="bold",
                    bbox=dict(boxstyle="round,pad=0.12", fc="white", ec="none", alpha=0.9), zorder=5)
    obs_basins: dict = {}
    for lab, rel, d25, d26, _ in PAIRS:
        for d in (d25, d26):
            obs_basins.setdefault((rel, d), set()).update(PAIR_BASINS[lab])
    for _, r in rain.iterrows():
        d, rel = str(r["date"]), int(r["rel"])
        yr = d[:4]
        if (rel, yr) not in lane_y:
            continue
        u = (rel, d) in used
        bs = obs_basins.get((rel, d), BASINS5)
        vals = [r[b] for b in bs if pd.notna(r[b])]
        mx = max(vals) if vals else np.nan
        wet = (mx >= 50) if u else (isinstance(r["wet"], str) and r["wet"].strip() != "")
        ax.scatter([X(d)], [lane_y[(rel, yr)]], s=150 if u else 60, marker=SAT_MK[r["sat"]],
                   facecolor="#6baed6" if wet else "white", edgecolor=C_INK if u else "#9a9a9a",
                   lw=1.6 if u else 0.9, zorder=3)
        if u and np.isfinite(mx):
            ax.text(X(d), lane_y[(rel, yr)] + 0.30, f"{mx:.0f}", ha="center", va="bottom", fontsize=FS - 6,
                    color="#2171b5" if wet else C_MUTED, zorder=4)
    ax.set_yticks([lane_y[k] for k in lane_y])
    ax.set_yticklabels([f"{lab}  {yr}" for rel, lab in LANES for yr in ("2025", "2026")], fontsize=FS - 2)
    ax.tick_params(axis="y", length=0)
    ax.set_xticks([0, 14, 31, 45, 61]); ax.set_xticklabels(["7/1", "7/15", "8/1", "8/15", "8/31"])
    ax.set_xlim(-1, 62); ax.set_ylim(lane_y[(127, "2026")] - 0.7, 0.9)
    for s_ in ("top", "right", "left"):
        ax.spines[s_].set_visible(False)
    ax.grid(axis="x", color=C_GRID, lw=0.8); ax.set_axisbelow(True)
    ax.set_xlabel("촬영일 (UTC, 두 해를 같은 달력에 겹침) · 숫자 = 대상 권역 직전 5일 강우(권역 중앙값, mm)", color=C_MUTED)
    ax.legend(handles=[Line2D([], [], color=C_AUG, lw=3), Line2D([], [], color=C_NAVY, lw=1.3, ls=(0, (4, 2))),
                       Line2D([], [], color="#555555", lw=1.6, ls=(0, (1, 1.5))),
                       Line2D([], [], marker="o", ls="", ms=10, mfc="white", mec=C_INK),
                       Line2D([], [], marker="s", ls="", ms=9, mfc="white", mec=C_INK),
                       Line2D([], [], marker="^", ls="", ms=10, mfc="white", mec=C_INK),
                       Line2D([], [], marker="o", ls="", ms=10, mfc="#6baed6", mec=C_INK)],
              labels=["본쌍 (7월 S1C↔S1C · 8월 S1C↔S1D)", "민감도 4쌍 (S1C↔S1C)", "건조 대조 (저수지)",
                      "S1C", "S1D", "S1A", "직전 5일 ≥50 mm"],
              loc="lower left", bbox_to_anchor=(-0.12, 1.02), ncol=4, fontsize=FS - 4, handlelength=1.8,
              columnspacing=1.0, handletextpad=0.4)
    fig.savefig(out, facecolor="white"); plt.close(fig)


def fig7(cn: pd.DataFrame, out: Path) -> None:
    sw = pd.read_csv(WAMIS_DIR / "sar_vs_wamis_v11.csv", encoding="utf-8-sig")
    dry = pd.read_csv(GEE / "watershed_change_v11_aug_a_dry0812.csv", encoding="utf-8-sig").set_index("대상")
    dams = ["안동", "임하", "영천", "성덕", "밀양", "운문"]
    wet12 = {"밀양", "운문"}                                       # 0812 도 54~82 mm — 건조 대조가 아니다
    fig = plt.figure(figsize=(14.8, 5.9))
    gs = fig.add_gridspec(1, 2, wspace=0.28, left=0.09, right=0.985, top=0.84, bottom=0.15)
    panels = []
    for kind in ("res", "bas"):
        rows = []
        for d in dams:
            if kind == "res":
                s = sw[sw.저수지 == RES8[d]].set_index("월")["SAR_dA_pct"]
                rows.append((d, s["8월"], s["8월 0812→0802 건조대조"]))
            else:
                rows.append((d, cget(cn, d, "8월").pct, float(dry.loc[d, "Δ%"])))
        panels.append(rows)
    titles = ["(a) 저수지 폴리곤 — 차 ≤1.0%p", "(b) 댐유역 — 최대 6.2%p (안동·임하)"]
    for k, rows in enumerate(panels):
        ax = fig.add_subplot(gs[0, k])
        for i, (d, a, b) in enumerate(rows):
            yi = len(rows) - 1 - i
            grey = k == 1 and d in wet12
            ax.plot([a, b], [yi, yi], color="#c9c9c9", lw=2.4, zorder=1)
            ax.scatter([a], [yi], s=120, marker="s", color="#cccccc" if grey else C_AUG, edgecolor="white", lw=1.3,
                       zorder=3)
            ax.scatter([b], [yi], s=95, marker="o", facecolor="white", edgecolor="#bbbbbb" if grey else C_NAVY,
                       lw=1.6, zorder=4)
            ax.text(max(a, b) + 1.0, yi, f"{f1(b - a)}%p" + ("  (0812도 젖음)" if grey else ""), va="center",
                    fontsize=FS - 4, color=C_MUTED if grey else C_INK)
        ax.set_yticks(range(len(rows)))
        ax.set_yticklabels([RES8[d] if k == 0 else d for d, _, _ in rows][::-1], fontsize=FS - 2)
        ax.set_xlim(-30, 8)
        ax.set_xticks([-30, -20, -10, 0])
        ax.set_xticklabels([str(v).replace("-", M) for v in (-30, -20, -10, 0)])
        ax.set_xlabel("SAR 수체 면적 Δ% (8월A, 2026 = 08-02 고정)", color=C_MUTED)
        ax.grid(axis="x", color=C_GRID, lw=0.8); ax.set_axisbelow(True)
        for s_ in ("top", "right", "left"):
            ax.spines[s_].set_visible(False)
        ax.tick_params(axis="y", length=0)
        ax.set_title(titles[k], loc="left", fontweight="bold")
    fig.legend(handles=[Line2D([], [], marker="s", ls="", ms=11, color=C_AUG),
                        Line2D([], [], marker="o", ls="", ms=10, mfc="white", mec=C_NAVY, mew=1.6)],
               labels=["본쌍 2025-08-06 (직전 5일 낙동강 102 mm)", "건조 대조 2025-08-12 (33 mm)"],
               loc="upper center", bbox_to_anchor=(0.5, 1.0), ncol=2, fontsize=FS - 2)
    fig.savefig(out, facecolor="white"); plt.close(fig)


# ════════════════════════ 표 — 그리기·내보내기 ════════════════════════
@dataclass
class Tbl:
    key: str
    caption: str
    header: list
    rows: list                        # list[list[str]] | str(구역 제목 줄)
    notes: list = field(default_factory=list)
    align: list | None = None         # 'l' | 'c'
    cell: dict = field(default_factory=dict)   # (r, c) → {'bold','color','bg'};  (r, None) = 그 줄 전체
    sep: list = field(default_factory=list)     # 이 줄 뒤에 굵은 선


def _style(t: Tbl, r: int, c: int) -> dict:
    st = dict(t.cell.get((r, None), {}))
    st.update(t.cell.get((r, c), {}))
    return st


_MFIG = None


def _tw(s: str, fs: float, bold: bool = False) -> float:
    global _MFIG
    if _MFIG is None:
        _MFIG = plt.figure(dpi=72)
    t = _MFIG.text(0, 0, s, fontsize=fs, fontweight="bold" if bold else "normal")
    w = t.get_window_extent(_MFIG.canvas.get_renderer()).width
    t.remove()
    return w


def _wrap(s: str, width: float, fs: float) -> list:
    out, cur = [], ""
    for tok in s.split(" "):
        cand = (cur + " " + tok) if cur else tok
        if _tw(cand, fs) <= width or not cur:
            cur = cand
        else:
            out.append(cur); cur = tok
    if cur:
        out.append(cur)
    return out


def render_png(t: Tbl, out: Path, fs: float = 11.0) -> list:
    hfs, pad, lh = fs + 0.5, 7.0, fs * 1.32
    ncol = len(t.header)
    al = t.align or (["l"] + ["c"] * (ncol - 1))
    widths = [_tw(h, hfs, True) for h in t.header]
    for r in t.rows:
        if isinstance(r, str):
            continue
        for c, v in enumerate(r):
            widths[c] = max(widths[c], _tw(str(v), fs, True))
    widths = [w + 2 * pad for w in widths]
    W = sum(widths)
    hh = max(h.count("\n") + 1 for h in t.header) * lh + 2 * pad
    rh = [((1 if isinstance(r, str) else max(str(v).count("\n") + 1 for v in r)) * lh + 1.3 * pad) for r in t.rows]
    cap = _wrap(t.caption, W, fs + 1.5)
    notes = [ln for n in t.notes for ln in _wrap(n, W, fs - 1.5)]
    capH = len(cap) * (fs + 1.5) * 1.35 + 6
    notesH = len(notes) * (fs - 1.5) * 1.38 + (8 if notes else 0)
    H = capH + hh + sum(rh) + notesH + 6
    m = 6.0
    fig = plt.figure(figsize=((W + 2 * m) / 72, (H + 2 * m) / 72))
    ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(-m, W + m); ax.set_ylim(H + m, -m); ax.axis("off")
    y = 0.0
    for ln in cap:
        ax.text(0, y, ln, fontsize=fs + 1.5, fontweight="bold", va="top", color=C_INK)
        y += (fs + 1.5) * 1.35
    y += 6
    ax.add_patch(Rectangle((0, y), W, hh, color=C_NAVY, lw=0))
    x = 0.0
    for c, h in enumerate(t.header):
        ax.text(x + widths[c] / 2, y + hh / 2, h, fontsize=hfs, fontweight="bold", color="white",
                ha="center", va="center", linespacing=1.15)
        x += widths[c]
    y += hh
    body = 0
    for i, r in enumerate(t.rows):
        h = rh[i]
        if isinstance(r, str):
            st = _style(t, i, 0)
            ax.add_patch(Rectangle((0, y), W, h, color=st.get("bg", "#e4e9f1"), lw=0))
            ax.text(pad, y + h / 2, r, fontsize=fs, fontweight="bold", color=st.get("color", C_NAVY), va="center")
            body = 0
        else:
            bg = "#f4f7fb" if body % 2 else "white"
            st_row = t.cell.get((i, None), {})
            ax.add_patch(Rectangle((0, y), W, h, color=st_row.get("bg", bg), lw=0))
            x = 0.0
            for c, v in enumerate(r):
                st = _style(t, i, c)
                if "bg" in st and (i, c) in t.cell:
                    ax.add_patch(Rectangle((x, y), widths[c], h, color=st["bg"], lw=0))
                xa = x + (pad if al[c] == "l" else widths[c] / 2)
                ax.text(xa, y + h / 2, str(v), fontsize=fs, ha="left" if al[c] == "l" else "center", va="center",
                        fontweight="bold" if st.get("bold") else "normal", color=st.get("color", C_INK),
                        linespacing=1.15)
                x += widths[c]
            body += 1
        y += h
        ax.plot([0, W], [y, y], color="#7a8796" if i in t.sep else "#d5dbe3", lw=1.3 if i in t.sep else 0.6)
    ax.plot([0, W], [y, y], color=C_NAVY, lw=1.4)
    y += 8
    for ln in notes:
        ax.text(0, y, ln, fontsize=fs - 1.5, va="top", color=C_MUTED)
        y += (fs - 1.5) * 1.38
    fig.savefig(out, dpi=300, facecolor="white"); plt.close(fig)
    return widths


def write_xlsx(tbls: list, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook(); wb.remove(wb.active)
    thin = Side(style="thin", color="D5DBE3")
    for t in tbls:
        ws = wb.create_sheet(t.key.split("_")[0])
        ncol = len(t.header)
        al = t.align or (["l"] + ["c"] * (ncol - 1))
        ws.cell(1, 1, t.caption).font = Font(bold=True, size=12)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
        for c, h in enumerate(t.header, 1):
            cl = ws.cell(2, c, h.replace("\n", " "))
            cl.font = Font(bold=True, color="FFFFFF"); cl.fill = PatternFill("solid", fgColor="00478B")
            cl.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, r in enumerate(t.rows):
            rr = 3 + i
            if isinstance(r, str):
                cl = ws.cell(rr, 1, r); cl.font = Font(bold=True, color="00478B")
                cl.fill = PatternFill("solid", fgColor="E4E9F1")
                ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=ncol)
                continue
            for c, v in enumerate(r, 1):
                st = _style(t, i, c - 1)
                cl = ws.cell(rr, c, str(v).replace("\n", " "))
                cl.font = Font(bold=bool(st.get("bold")), color=st.get("color", "#222222").lstrip("#").upper())
                if "bg" in st:
                    cl.fill = PatternFill("solid", fgColor=st["bg"].lstrip("#").upper())
                cl.alignment = Alignment(horizontal="left" if al[c - 1] == "l" else "center",
                                         vertical="center", wrap_text=True)
                cl.border = Border(bottom=thin)
        rr = 3 + len(t.rows) + 1
        for n in t.notes:
            ws.cell(rr, 1, n).font = Font(size=9, color="666666"); rr += 1
        for c in range(1, ncol + 1):
            col = [len(t.header[c - 1])] + [len(str(r[c - 1])) for r in t.rows if not isinstance(r, str)]
            ws.column_dimensions[ws.cell(2, c).column_letter].width = min(46, max(8, 1.6 * max(col) + 2))
        ws.freeze_panes = "A3"
    wb.save(path)


def write_pptx(tbls: list, widths_png: dict, path: Path) -> None:
    """포스터와 같은 A0 세로 · 머리 22 pt · 본문 20 pt 굵게 · 맑은 고딕. 슬라이드당 표 하나."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Cm, Emu, Pt

    def setfont(run, size, bold, color):
        f = run.font
        f.size, f.bold, f.name = Pt(size), bold, "맑은 고딕"
        f.color.rgb = RGBColor.from_string(color.lstrip("#").upper())
        rpr = run._r.get_or_add_rPr()
        ea = rpr.find(qn("a:ea"))
        if ea is None:
            ea = rpr.makeelement(qn("a:ea"), {}); rpr.append(ea)
        ea.set("typeface", "맑은 고딕")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(84.1), Cm(118.9)
    blank = prs.slide_layouts[6]
    scale = 20.0 / 11.0                                 # PNG 본문 11 pt → pptx 20 pt
    for t in tbls:
        s = prs.slides.add_slide(blank)
        left, top = Cm(4.0), Cm(4.0)
        wpt = [w * scale for w in widths_png[t.key]]
        Wtot = int(sum(wpt) * 12700)
        tb = s.shapes.add_textbox(left, top, Wtot, Cm(1.6))
        tf = tb.text_frame; tf.word_wrap = True
        r0 = tf.paragraphs[0].add_run(); r0.text = t.caption
        setfont(r0, 22, True, C_INK)
        nrow = 1 + len(t.rows)
        tab = s.shapes.add_table(nrow, len(t.header), left, top + Cm(2.2), Wtot, Cm(1.3 * nrow)).table
        tab.first_row = True
        for c, w in enumerate(wpt):
            tab.columns[c].width = Emu(int(w * 12700))
        al = t.align or (["l"] + ["c"] * (len(t.header) - 1))

        def fill(cell, text, size, bold, color, bg, align):
            cell.margin_left = cell.margin_right = Cm(0.15)
            cell.margin_top = cell.margin_bottom = Cm(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(bg.lstrip("#").upper())
            tf = cell.text_frame; tf.word_wrap = True
            for k, ln in enumerate(str(text).split("\n")):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.LEFT if align == "l" else PP_ALIGN.CENTER
                r = p.add_run(); r.text = ln
                setfont(r, size, bold, color)

        for c, h in enumerate(t.header):
            fill(tab.cell(0, c), h, 22, True, "#FFFFFF", C_NAVY, "c")
        body = 0
        for i, r in enumerate(t.rows):
            rr = i + 1
            if isinstance(r, str):
                st = _style(t, i, 0)
                a = tab.cell(rr, 0)
                a.merge(tab.cell(rr, len(t.header) - 1))
                fill(a, r, 20, True, st.get("color", C_NAVY), st.get("bg", "#E4E9F1"), "l")
                body = 0
                continue
            row_bg = t.cell.get((i, None), {}).get("bg", "#F4F7FB" if body % 2 else "#FFFFFF")
            for c, v in enumerate(r):
                st = _style(t, i, c)
                fill(tab.cell(rr, c), v, 20, True, st.get("color", C_INK),
                     st["bg"] if ((i, c) in t.cell and "bg" in st) else row_bg, al[c])
            body += 1
        if t.notes:
            nt = s.shapes.add_textbox(left, top + Cm(2.4 + 1.25 * nrow), Wtot, Cm(1.0 * len(t.notes)))
            ntf = nt.text_frame; ntf.word_wrap = True
            for k, n in enumerate(t.notes):
                p = ntf.paragraphs[0] if k == 0 else ntf.add_paragraph()
                r = p.add_run(); r.text = n
                setfont(r, 16, False, "#555555")
    prs.save(path)


def md_table(t: Tbl) -> str:
    def esc(v):
        return str(v).replace("|", "\\|").replace("\n", "<br>")
    al = t.align or (["l"] + ["c"] * (len(t.header) - 1))
    out = ["| " + " | ".join(esc(h) for h in t.header) + " |",
           "|" + "|".join(":--" if a == "l" else ":-:" for a in al) + "|"]
    for i, r in enumerate(t.rows):
        if isinstance(r, str):
            out.append("| **" + esc(r) + "** |" + " |" * (len(t.header) - 1))
            continue
        cells = []
        for c, v in enumerate(r):
            st = _style(t, i, c)
            s = esc(v)
            if st.get("bold") and s and s != "—":
                s = f"**{s}**"
            cells.append(s)
        out.append("| " + " | ".join(cells) + " |")
    return "\n".join(out)


# ════════════════════════ 표 내용 ════════════════════════
HI = {"bold": True, "color": C_HI}
B = {"bold": True}
GREY = {"color": "#8a8a8a"}
UP = {"color": C_UP, "bold": True}
SUM = {"bg": "#fff1e6", "bold": True}


def kst(t) -> str:
    """'2025072518' → '07-25 18시' (WAMIS 시자료는 KST)."""
    t = str(t)
    return f"{t[4:6]}-{t[6:8]} {int(t[8:10])}시"


def lvl(st: dict | None) -> str:
    return "—" if st is None else f"{st['val']:.2f} m ({f2(st['anom'], True)})"


def build_tables(cn, mon, sw, swsens, sens) -> list:
    T = []
    rk = pd.read_csv(RANK_CSV, encoding="utf-8-sig", dtype={"코드": str})

    # 표1 ─ 본문 8개
    rows, cell = [], {}
    for i, d in enumerate(DAMS8):
        j, a = cget(cn, d, "7월"), cget(cn, d, "8월")
        st = wstat(mon, d)
        sv = sw[(sw.저수지 == RES8[d]) & (sw.월 == "8월")].iloc[0]
        dk = SMALL_DKM2.get(d, a.w26 - a.w25)
        rows.append([REGION8[d], f"{d} ({RES8[d]})", STAGE[d], lvl(st), where(st),
                     anom_s(st, "a25") + (" ⓒ" if d == "안동" else ""),
                     "제외 ⓕ" if d == "평림" else f1(j.pct), f1(a.pct),
                     f"{a.w25:.2f} → {a.w26:.2f} ({f2(dk, True)})" + ("†" if d in SMALL_DKM2 else ""),
                     f1(sv.SAR_dA_pct), f1(sv.WAMIS_dV_pct), f"{sv.k_eff:.2f}"])
        cell[(i, 7)] = HI
        if STAGE[d] in ("심각", "경계"):
            cell[(i, 2)] = B
        if where(st).endswith("최저"):
            cell[(i, 4)] = B
        if d == "평림":
            cell[(i, 6)] = GREY
    sts = [wstat(mon, d) for d in DAMS8]
    n_below = sum(s["anom"] < 0 for s in sts)
    n_low = sum(not s["lower"] for s in sts)
    n_after22 = sum(s["lower"] == [2022] for s in sts)
    n_up25 = sum(s["a25"] > 0 for s in sts)
    rows.append(["요약", "8곳", "8/8 공식 단계", f"{n_below}/8 중앙값 아래", f"{n_low}곳 10년 최저 · {n_after22}곳 2022년 다음",
                 f"{n_up25}/8 중앙값 위", f"{M}12.7 ~ {M}26.7 (7/8)", f"{M}4.7 ~ {M}27.7 (8/8)", "", "", "", "8/8 > 1"])
    cell[(8, None)] = SUM
    T.append(Tbl("표1_본문8개댐유역_핵심결과",
                 "Table 1. 본문 8개 댐유역 — 정부 예·경보, WAMIS 댐별 10년 기록, SAR 수체 면적 변화(2025→2026)",
                 ["권역", "댐유역\n(저수지)", "8월 가뭄\n예·경보", "WAMIS 8월 저수위 2026\n(10년 중앙값 대비 m)",
                  "10년 기록에서\n2026-08", "2025-08\n(중앙값 대비 m)", "SAR 유역\n7월 Δ%", "SAR 유역\n8월 Δ%",
                  "8월 유역 수체 km²\n2025 → 2026 (Δkm²)", "저수지 SAR\n8월 Δ%", "WAMIS 저수량\n8월 Δ%",
                  "k_eff\n(저수량÷면적)"],
                 rows,
                 notes=["8월 가뭄 예·경보 = 행정안전부 「8월 가뭄 예·경보」(2026-08-10 발표)의 댐 단계 — 예년 대비 저수량에 따른 정부 판단. "
                        "WAMIS = 각 댐의 8월 월평균 저수위를 그 댐의 2017~2026년 10개 값과 비교(다른 댐과 견주지 않음).",
                        "2025 열: 전년도가 자기 10년 중앙값보다 높았던 댐은 전년 대비 Δ% 가 평년 대비보다 크게 보인다. ⓒ 안동은 2025 도 "
                        "중앙값보다 낮아 Δ% 가 작게 보인다(2026 은 10년 최저) — 「덜 나쁘다」로 읽지 않는다.",
                        "SAR 쌍: 7월 낙동강 6곳 ASC54 0725↔0715, 섬진강 DESC134 0718↔0720(모두 S1C↔S1C) · 8월 낙동강 6곳 ASC54 "
                        "0806↔0802, 섬진강·평림 ASC127 0811↔0807(S1C↔S1D — 2025 년에 S1D 가 없어 불가피, 두 위성 면적 차 중앙 0.5%).",
                        "k_eff = 저수량 Δ% ÷ 저수지 SAR 면적 Δ%, 1 보다 크면 SAR 감소가 저수량 감소보다 작다(하한). "
                        "† 수체 1 km² 미만 — Δ% 와 Δkm² 를 함께 읽는다. ⓕ 평림 7월(−1.5%)은 두 해 임계가 3구간 달라 같은 임계로 맞추면 "
                        "+0.2~+0.7% → 감소로 세지 않는다."],
                 align=["c", "l"] + ["c"] * 10, cell=cell, sep=[7]))

    # 표2 ─ 남부 WAMIS 댐 전체, 자기 10년 기록
    rows, cell = [], {}
    dams = (mon[["권역", "댐"]].drop_duplicates().assign(o=lambda d: d["권역"].map(WREG_ORDER.index))
            .sort_values(["o"], kind="stable"))
    cur = None
    stats = []
    for _, r in dams.iterrows():
        if r["권역"] != cur:
            cur = r["권역"]; rows.append(cur.replace(" · ", "·"))
        st, s7 = wstat(mon, r["댐"]), wstat(mon, r["댐"], 7)
        stats.append(st)
        sar_nm = {"주암(본)": "주암", "김천부항": "부항"}.get(r["댐"], r["댐"])
        stg = STAGE.get(sar_nm, "—")
        rows.append([r["댐"], st["mgr"], f"{st['n']}년", f"{st['val']:.2f}", f2(st["anom"], True), where(st),
                     f2(s7["anom"], True), f2(st["a25"], True), stg])
        i = len(rows) - 1
        if where(st).endswith("최저"):
            cell[(i, 5)] = HI
        if st["anom"] > 0.005:
            cell[(i, 4)] = UP; cell[(i, 5)] = UP
        if stg in ("심각", "경계"):
            cell[(i, 8)] = B
    nb = sum(s["anom"] < -0.005 for s in stats)
    nl = sum(not s["lower"] for s in stats)
    nl6 = sum((not s["lower"]) and s["n"] < 10 for s in stats)
    rows.append(["요약", "", "", "", f"{nb}/{len(stats)} 아래", f"최저 {nl}곳 (6년 기록 {nl6})", "", "", "8곳 단계"])
    cell[(len(rows) - 1, None)] = SUM
    T.append(Tbl("표2_남부WAMIS댐_10년기록",
                 "Table 2. 남부 WAMIS 댐 2026년 저수위 — 댐마다 자기 10년(2017~2026) 기록과 비교",
                 ["댐", "관리", "기록", "2026-08 저수위\n(EL.m)", "8월 10년 중앙값\n대비 (m)", "10년 기록에서\n2026-08",
                  "7월 중앙값\n대비 (m)", "2025-08 중앙값\n대비 (m)", "8월 가뭄\n예·경보"],
                 rows,
                 notes=["WAMIS 월자료(mn_mndata) 월평균 저수위. 조정지 댐 제외. API 가 2017 이전 월자료를 주지 않아 기록은 10년(농어촌공사 4곳은 "
                        "6년, 2021~). 중앙값은 그 댐 기록의 표준 중앙값. 「○○년 다음으로 낮음」 = 2026 보다 낮았던 해가 그 해들뿐.",
                        "2026 보다 낮았던 해는 대부분 2017·2022년이다. 8월 가뭄 예·경보 = 행안부 8.10 발표에 단계가 있는 댐(— = 단계 없음). "
                        "한수원·농어촌공사 댐은 발전·관개 운영 수위라 저수위 신호가 약할 수 있다."],
                 align=["l", "c", "c", "c", "c", "l", "c", "c", "c"], cell=cell))

    # 표3 ─ 남부 15 대권역 SAR
    rows, cell = [], {}
    by_reg = {}
    for st in stats:
        by_reg.setdefault(st["region"], []).append(st)
    remark = {"낙동강": "보도자료: 유역 댐 저수량 예년 이하", "섬진강": "보도자료: 유역 댐 저수량 예년 이하",
              "형산강": "안계 2026-08 중앙값 위", "회야·수영": "선암 2026-08 중앙값 위",
              "만경·동진": "하구·간척지 — 저수지와 다른 물 ⓔ", "금강": "용담 10년 최저 · 대청 −0.54 m"}
    for lab, bbn, wk in REG:
        j7, a8 = cget(cn, bbn, "7월", "대권역"), cget(cn, bbn, "8월", "대권역")
        g = by_reg.get(wk, [])
        rows.append([lab, f1(j7.pct), f1(a8.pct), f"{a8.w25:.2f} → {a8.w26:.2f}",
                     f"{sum(s['anom'] < -0.005 for s in g)}/{len(g)}", remark.get(lab, "")])
        if a8.pct > 0:
            cell[(len(rows) - 1, 2)] = UP
    rows.append("WAMIS 댐이 없는 대권역")
    for lab in NO_DAM:
        j7, a8 = cget(cn, lab, "7월", "대권역"), cget(cn, lab, "8월", "대권역")
        rows.append([lab, f1(j7.pct), f1(a8.pct), f"{a8.w25:.2f} → {a8.w26:.2f}", "—", ""])
        if a8.pct > 0:
            cell[(len(rows) - 1, 2)] = UP
    T.append(Tbl("표3_남부15대권역_SAR",
                 "Table 3. 남부 15개 대권역 — SAR 수체 면적 Δ%(2025→2026)와 권역 안 WAMIS 댐의 2026년 8월 저수위",
                 ["대권역", "SAR 7월\nΔ%", "SAR 8월\nΔ%", "8월 수체 km²\n2025 → 2026",
                  "WAMIS 댐 중 2026-08\n자기 중앙값 아래", "비고"],
                 rows,
                 notes=["SAR 쌍(정본 배정): 7월 낙동강·동남권 ASC54 0725↔0715 / 금강·섬진·영산 계열 DESC134 0718↔0720, "
                        "8월 ASC54 0806↔0802 / ASC127 0811↔0807. SAR Δ% 는 판정이 아니라 관측(하한)이다.",
                        "WAMIS 열은 권역 안 각 댐을 자기 10년 기록과 비교한 결과를 센 것이다(댐별 값은 Table 2). 권역을 순위로 판정하지 않는다. "
                        "정부 보도자료는 「낙동강·섬진강 유역 다목적댐 11곳·용수댐 10곳 저수량 예년 이하」라고 판단했다."],
                 align=["l", "c", "c", "c", "c", "l"], cell=cell, sep=[3]))

    # 표4 ─ 정부 예·경보
    order = ["운문", "밀양", "섬진강", "안동", "임하", "성덕", "영천", "평림"]
    rows, cell = [], {}
    for i, d in enumerate(order):
        st = wstat(mon, d)
        rows.append([f"{d}댐", REGION8[d], STAGE[d], STAGE_ENTRY.get(d, "—"), lvl(st), where(st),
                     f1(cget(cn, d, "8월").pct)])
        if STAGE[d] in ("심각", "경계"):
            cell[(i, 2)] = HI
        if where(st).endswith("최저"):
            cell[(i, 5)] = B
    rows.append("대조 — 금강: WAMIS 기록과 보도자료가 다르게 말한다")
    for d in ("용담", "대청"):
        st = wstat(mon, d)
        rows.append([f"{d}댐", "금강", "단계 없음", "—", lvl(st), where(st),
                     f1(cget(cn, d, "8월").pct) + (" (변화 없음)" if d == "대청" else "")])
        cell[(len(rows) - 1, None)] = {"color": "#555555"}
    T.append(Tbl("표4_공식가뭄예경보_대조",
                 "Table 4. 정부 가뭄 예·경보와 WAMIS 댐별 10년 기록 — 본문 8개 댐은 모두 8월 공식 가뭄 단계였다",
                 ["댐", "권역", "8월 예·경보 단계\n(8.10 발표)", "현 단계 진입일\n(2026, 기후부)",
                  "WAMIS 8월 저수위 2026\n(10년 중앙값 대비 m)", "10년 기록에서\n2026-08", "SAR 유역\n8월 Δ%"],
                 rows,
                 notes=["보도자료(행안부 「8월 가뭄 예·경보」 2026-08-10, 8월 1일 기준): 다목적댐 19곳 저수량 예년의 104.3%, 용수댐 12곳 66.1% — "
                        "전국으로는 안정적이나 낙동강·섬진강 유역 다목적댐 11곳·용수댐 10곳이 예년 이하. 6개월 누적 강수 평년(1991~2020)의 82.3%, "
                        "기상가뭄 48개 시·군. 농업용 저수지 53.2%(평년 69.9% 의 76.1%).",
                        "진입일 = 기후에너지환경부 보도참고자료 원문(운문 심각 7.15 · 안동·임하 주의 7.8 · 성덕 주의 7.18 · 밀양 경계 8.2 · "
                        "평림 주의 8.2). 영천·섬진강은 진입 보도자료 미확보.",
                        "금강: 용담은 WAMIS 10년 최저인데 공식 단계가 없다. 공식 단계는 보도자료의 예년 저수량 기준 판단이고 WAMIS 는 댐별 10년 "
                        "수위 기록이라 판단 기준이 다르다 — 각자의 기준으로 따로 적는다. 대청 SAR Δ% 는 궤도에 따라 −4.4 ↔ +0.5% 로 갈려 인용하지 않는다."],
                 align=["l", "c", "c", "c", "c", "l", "c"], cell=cell, sep=[7]))

    # 표5 ─ SAR 면적 ↔ WAMIS 저수량
    rows, cell = [], {}
    sw2 = sw[sw.월.isin(["7월", "8월"])]
    for d in DAMS8:
        res = RES8[d]
        for mth in ("7월", "8월"):
            r = sw2[(sw2.저수지 == res) & (sw2.월 == mth)].iloc[0]
            rows.append([res if mth == "7월" else "", mth, f"{kst(r.t0)} → {kst(r.t1)}", f"{r.rsrt0:.1f} → {r.rsrt1:.1f}",
                         f2(r.dh_m, True), f1(r.WAMIS_dV_pct), f1(r.SAR_dA_pct), f"{r.k_eff:.2f}",
                         "하한" if r.k_eff > 1 else "—"])
            cell[(len(rows) - 1, 6)] = B
    nmain = sum(1 for r in rows if r[-1] == "하한")
    k = swsens["k_eff"]
    rows.append(["집계", "", "본쌍 16", "", "", "", "", f"{sw2.k_eff.min():.2f}~{sw2.k_eff.max():.2f}", f"{nmain}/16"])
    rows.append(["", "", "민감도 2×2 (16칸, Table 6)", "", "", "", "", f"{k.min():.2f}~{k.max():.2f}",
                 f"{int((k > 1.0).sum())}/16 (동률 1·역전 1)"])
    rows.append(["", "", "합계 32칸", "", "", "", "", "", f"{nmain + int((k > 1.0).sum())}/32"])
    for r_ in range(len(rows) - 3, len(rows)):
        cell[(r_, None)] = SUM
    T.append(Tbl("표5_SAR면적_WAMIS저수량_하한",
                 "Table 5. 저수지 SAR 면적 Δ% ↔ WAMIS 저수량 Δ% — 같은 1:1 쌍·같은 촬영 시각, 32칸 중 30칸에서 SAR 감소가 더 작다(하한)",
                 ["저수지", "월", "촬영 시각 (KST)\n2025 → 2026", "저수율 (%)\n2025 → 2026", "수위 Δ\n(m)", "저수량\nΔ%",
                  "SAR 면적\nΔ%", "k_eff", "판정"],
                 rows,
                 notes=["k_eff = 저수량 Δ% ÷ SAR 면적 Δ%. 저수지는 위로 갈수록 넓어 수위가 내려가면 부피가 면적보다 빨리 준다 → "
                        "형상만으로 k 1.3~2.0 이 기대되나 16쌍 중 14쌍은 그 범위로도 설명되지 않는다(안동 8월 1.62·옥정 7월 1.89 만 범위 안).",
                        "하한이 깨진 2칸은 모두 안동·2025-08-23 기준 민감도 쌍(k 0.99 동률, 0.77 역전): SAR 은 네 쌍에서 −11.5~−12.2 로 "
                        "거의 고정인데 저수량이 −9.6~−16.7 로 날짜를 따라 움직인다(분모 쪽). 사후 점검(저수량 월평균)으로는 16/16 하한 — 본판 아님.",
                        "7월 옥정·평림은 DESC134 0718↔0720(KST 익일 06시). 저수지 SAR 은 gee reservoir_km2_v11 의 날짜 열(궤도혼합 열 아님)."],
                 align=["l", "c", "c", "c", "c", "c", "c", "c", "c"], cell=cell, sep=[15]))

    # 표6 ─ 민감도 2×2
    pairs = ["0811→0813", "0823→0813", "0811→0825", "0823→0825"]
    rain = {"안동호": "4.6 · 0.7 | 40.1 · 36.6", "임하호": "7.7 · 0.0 | 18.8 · 35.2", "옥정호": "43.0 · 0.0 | 0.0 · 4.5",
            "평림호": "57.5 · 0.0 | 0.0 · 0.0", "장성호": "—"}
    main_orb = {"안동호": "ASC54", "임하호": "ASC54", "옥정호": "ASC127", "평림호": "ASC127", "장성호": "ASC127"}
    rows, cell = [], {}
    for i, res in enumerate(["안동호", "임하호", "옥정호", "평림호", "장성호"]):
        cells = []
        for j, p in enumerate(pairs):
            kk = swsens[(swsens.저수지 == res) & swsens.월.str.endswith(p)]
            cells.append(f1(sens.loc[res, p]) + (f"\n(k {kk.k_eff.iloc[0]:.2f})" if len(kk) else ""))
            if len(kk) and kk.k_eff.iloc[0] <= 1.0:
                cell[(i, 2 + j)] = HI
        v4 = sens.loc[res, pairs].astype(float).round(1)      # 표에 보이는 값끼리의 폭 (근거 문서와 같은 방식)
        rows.append([res + ("*" if res == "장성호" else ""), f"{f1(sens.loc[res, 'main'])}\n({main_orb[res]})", *cells,
                     f"{v4.max() - v4.min():.1f}", rain[res]])
        cell[(i, 1)] = B
    neg = int((sens[pairs] < 0).to_numpy().sum())
    least = int((sens["main"] > sens[pairs].max(axis=1)).sum())
    rows.append(["요약", f"{least}/5 감소 최소", f"{neg}/20 감소", "", "", "", "", ""])
    cell[(5, None)] = SUM
    T.append(Tbl("표6_민감도2x2_DESC134",
                 "Table 6. 8월 민감도 4쌍(DESC134, S1C↔S1C) — 위성·날짜·강우를 바꿔도 20칸 모두 감소, 본쌍이 가장 보수적",
                 ["저수지", "본쌍 8월\n(S1C↔S1D)", "2025-08-11\n→ 2026-08-13", "2025-08-23\n→ 2026-08-13",
                  "2025-08-11\n→ 2026-08-25", "2025-08-23\n→ 2026-08-25", "네 쌍 폭\n(%p)",
                  "직전 5일 강우 (mm)\n2025 0811·0823 | 2026 0813·0825"],
                 rows,
                 notes=["SAR 저수지 면적 Δ%(괄호 k_eff = 같은 시각 WAMIS 저수량 Δ% ÷ SAR Δ%). 범위는 DESC134 가 네 날짜 모두 덮는 저수지 — "
                        "밀양·영천·성덕·운문은 미커버. *장성호(영산강, 농어촌공사)는 8개 댐 밖 대조 저수지(WAMIS 저수량 대조 없음).",
                        "본쌍이 5/5 곳에서 네 쌍 범위 밖 작은 쪽 → 보고 수치는 과대가 아니라 과소평가. 크기는 옥정에서 흔들린다(폭 5.5%p, "
                        "거의 전부 2026 쪽 날짜 0813→0825 — 두 날 다 건조). 붉은 칸 = 하한이 깨진 안동 2칸."],
                 align=["l"] + ["c"] * 7, cell=cell, sep=[4]))

    # 표7 ─ S2 교차검증
    s1t = pd.read_csv(S2_DIR / "표1_16사례_S2_SARv11.csv", encoding="utf-8-sig")
    s2t = pd.read_csv(S2_DIR / "표2_화소대응_v11.csv", encoding="utf-8-sig", dtype={"S2 2025": str, "S2 2026": str})
    s2t = s2t.dropna(subset=["댐"])
    s2t = s2t[s2t["댐"].isin(["영천", "성덕", "운문", "밀양", "안동", "임하"])].set_index("댐")
    rows, cell = [], {}
    for i, d in enumerate(["운문", "영천", "밀양", "성덕", "안동", "임하"]):
        a = s1t[s1t["사례"] == f"{d} 8월"].iloc[0]
        b = s2t.loc[d]
        dmn = float(a["ΔMNDWI"])
        d25, d26 = str(b["S2 2025"])[:8], str(b["S2 2026"])[:8]
        rows.append([d, STAGE[d], f"{d25[4:6]}-{d25[6:8]} → {d26[4:6]}-{d26[6:8]}" + ("*" if d == "밀양" else ""),
                     f"{float(a['공통유효화소%']):.1f}",
                     ("+" if dmn > 0 else "") + f"{dmn:.3f}".replace("-", M) + (" (≈0)" if d == "운문" else ""),
                     f1(float(b["SAR Δ%(저수지띠)"])), f1(float(b["S2 Δ%(저수지띠)"])),
                     (f"({float(b['SAR소실→S2소실 %']):.1f})†" if d == "성덕" else f"{float(b['SAR소실→S2소실 %']):.1f}"),
                     f"{float(b['2025수면 IoU']):.2f}", "본문" if float(a["공통유효화소%"]) >= 85 else "참고 (85% 미만)"])
        if d in ("운문", "영천", "밀양"):
            cell[(i, 7)] = HI
        if float(a["공통유효화소%"]) < 85:
            cell[(i, None)] = {"color": "#7a7a7a"}
    T.append(Tbl("표7_S2교차검증_화소대응",
                 "Table 7. Sentinel-2 교차검증(8월 ASC54 쌍) — 맑은 사례에서 두 센서가 같은 방향, SAR 소실 화소의 74~87% 를 광학도 소실로 판정",
                 ["댐", "8월\n예·경보", "S2 날짜\n2025 → 2026", "공통 유효\n화소 %", "유역평균\nΔMNDWI", "저수지띠\nSAR Δ%",
                  "저수지띠\nS2 Δ%", "SAR 소실 →\nS2 소실 %", "2025 수면\nIoU", "구분"],
                 rows,
                 notes=["SAR v11 2025-08-06↔2026-08-02. S2 는 L2A 단일일(합성 안 함). 유역평균 ΔMNDWI 는 정본 SCL 마스크, 저수지띠·화소 대응은 "
                        "수면 보존 마스크(SCL 2·7 유지 + 음수 반사도 0) — 두 지표를 한 문장에 섞지 않는다(유역평균은 유역의 2~4% 인 수면을 희석).",
                        "공통 유효화소 85% 이상 4사례: 수면 면적으로는 4/4 같은 방향, 유역평균으로는 3/4(운문 ≈0, 마스크 300 m 확장 시 −0.001). "
                        "85% 미만 12사례는 전부 ΔMNDWI 양수(잔여 대기효과). * 밀양 화소 대응은 S2 2026-08-06(08-01 은 수면 위 권운 71%), "
                        "유역평균은 08-01. † 성덕 화소 대응은 수체 0.66 km²·IoU 최저라 인용하지 않는다."],
                 align=["l"] + ["c"] * 9, cell=cell, sep=[3]))

    # 표8 ─ 관측쌍·전처리
    rainr = pd.read_csv(RAIN_CSV, dtype={"date": str})

    def rv(d, rel, basins):
        r = rainr[(rainr.date == d) & (rainr.rel == rel)].iloc[0]
        return " · ".join(f"{r[b]:.0f}" for b in basins)

    rows = [["7월A", "ASC54", "07-25 (S1C)", "07-15 (S1C)", "S1C↔S1C", "낙동강 6개 댐유역 · 낙동강·동남권 대권역",
             f"낙동강 {rv('20250725', 54, ['낙동강'])} → {rv('20260715', 54, ['낙동강'])}"],
            ["7월B", "DESC134", "07-18 (S1C)", "07-20 (S1C)", "S1C↔S1C", "섬진강·평림 · 금강·섬진강·영산강 계열",
             f"섬진·영산 {rv('20250718', 134, ['섬진강', '영산강'])} → {rv('20260720', 134, ['섬진강', '영산강'])}"],
            ["8월A", "ASC54", "08-06 (S1C)", "08-02 (S1D)", "S1C↔S1D", "낙동강 6개 댐유역 · 낙동강·동남권 대권역",
             f"낙동강 {rv('20250806', 54, ['낙동강'])} → {rv('20260802', 54, ['낙동강'])}"],
            ["8월B", "ASC127", "08-11 (S1C)", "08-07 (S1D)", "S1C↔S1D", "섬진강·평림 · 금강·섬진강·영산강 계열",
             f"섬진·영산 {rv('20250811', 127, ['섬진강', '영산강'])} → {rv('20260807', 127, ['섬진강', '영산강'])}"],
            "점검용 쌍 (본문 수치를 대체하지 않는다)",
            ["민감도 2×2", "DESC134", "08-11 · 08-23 (S1C)", "08-13 · 08-25 (S1C)", "S1C↔S1C", "안동·임하·옥정·평림 (+장성)",
             "Table 6"],
            ["건조 대조", "ASC54", "08-12 (S1A)", "08-02 (S1D)", "S1A↔S1D", "낙동강 6개 저수지 (8월A 대체)",
             f"낙동강 {rv('20250812', 54, ['낙동강'])} → {rv('20260802', 54, ['낙동강'])}"]]
    T.append(Tbl("표8_관측쌍_전처리",
                 "Table 8. 비교 설계 — 동일 상대궤도 1:1 관측쌍(전년도 2025 ↔ 2026, 월별)과 점검용 쌍",
                 ["쌍", "상대궤도", "2025 (위성)", "2026 (위성)", "위성", "대상", "직전 5일 강우 (mm)\n2025 → 2026"],
                 rows,
                 notes=["전처리(포스터 관측 13장 모두): Sentinel-1 GRD VH · 정밀궤도(POEORB) · 열잡음 제거·보정(β⁰) · Frost 3×3 · "
                        "Terrain-Flattening(γ⁰) · Range-Doppler 지형보정(COP30 DEM, EGM 지오이드 보정) · 10 m · EPSG:5179 · dB.",
                        "수체: 유역별 Otsu 임계(두 해 같은 구간 폭 0.1765 dB 히스토그램), 분리도 η 가 낮은 소유역은 상위 대권역 임계를 두 해 함께 상속. "
                        "8월은 2025 년에 S1D 가 없어 위성 교차가 불가피하다 — 같은 해·같은 궤도 두 위성 면적 차는 중앙 0.5%.",
                        "강우: WAMIS 댐유역 일강우, 촬영일 포함 직전 5일 합의 권역 중앙값(그 관측이 덮은 댐만). 강우는 거르는 기준이 아니라 해석용 주석이다."],
                 align=["l", "c", "c", "c", "c", "l", "c"], sep=[3]))

    # 표9 ─ 견고성 점검
    rows = [["위성 차이\n(8월 S1C↔S1D)", "같은 해·같은 궤도에서\n두 위성 면적 비교", "중앙 0.5%\n(유역 0.35%)",
             "8월 Δ% −4.7~−27.7% 를\n위성 차로 설명할 수 없다"],
            ["위성·날짜·강우\n(민감도 2×2)", "DESC134 S1C↔S1C 4쌍\n× 5개 저수지", "20/20 감소\n본쌍이 5/5 곳 가장 작은 감소",
             "보고 수치는 보수적\n(Table 6)"],
            ["직전 강우\n— 저수지", "2025 0806(102 mm) → 0812(33 mm)\n2026 0802 고정", "Δ% 차 ≤1.0%p\n6/6 하한 유지",
             "8월A 본쌍 감소율은\n0806 직전 비 때문이 아니다"],
            ["직전 강우\n— 댐유역", "같은 교체를\n유역 단위로", "최대 6.2%p(임하)·4.9%p(안동)\n담수호 안은 ≤1.05%",
             "차이는 담수호 밖 하천 — 젖은 날 하천이\n덜 잡힘 (안동·임하는 과소평가 쪽)"],
            ["임계 양자화\n(8월, ±1구간)", "0.1765 dB 를\n양쪽으로 직접 주입", "맑은 4유역 ≤1.10%p (상대 ≤4%)\n비대칭 0.00~0.20%p",
             "안동 1.00%p(13%)·임하 1.30%p(18%)\n— 8개 전체에 한 수로 쓰지 않는다"],
            ["임계 비대칭\n(7월)", "낙동강 두 해를\n같은 임계로 재산출", "6유역 모두 과소평가 쪽\n+0.3~+1.7%p",
             "7월 하한 −12.7(안동, 참값 −13.5~−13.7)은\n보수적 · 평림 7월 제외(ⓕ)"],
            ["SAR ↔ 저수량\n(하한)", "본쌍 16 + 민감도 16\n같은 촬영 시각", "30/32 하한\n안동 2칸 동률·역전",
             "「모두 하한」 이 아니라\n「32칸 중 30칸」"],
            ["커버리지", "유역 면적 중 관측 비율", "본문 8개 99.9~100%", "—"]]
    T.append(Tbl("표9_견고성점검_부록",
                 "Table 9. (부록) 견고성 점검 요약 — 무엇을 바꿔 봤고 결론이 어떻게 남았나",
                 ["점검", "방법", "결과", "해석 · 쓰는 범위"],
                 rows,
                 notes=["전부 실측(추정·외삽 없음). 임계 ±1구간은 「차(差) 항」만 덮고, 두 해 공통 오차는 민감도 4쌍이 따로 덮는다 — 섞지 않는다. "
                        "출처: gee DROUGHT_2026_EVIDENCE_KR §3-3·§3-4·§3-8·§4, S1 REVIEW_20261005 §2-1."],
                 align=["l", "l", "l", "l"]))

    # 표10 ─ 남부 댐유역 전수
    rows, cell = [], {}
    sdams = sorted(set(cn[cn.구분 == "댐유역"].대상))
    reg_of = {d: (wstat(mon, d)["region"] if wstat(mon, d) else NO_WAMIS_REG[d]) for d in sdams}
    sdams = sorted(sdams, key=lambda d: (WREG_ORDER.index(reg_of[d]), cget(cn, d, "8월").pct))
    cur = None
    for d in sdams:
        if reg_of[d] != cur:
            cur = reg_of[d]; rows.append(cur.replace(" · ", "·"))
        j, a = cget(cn, d, "7월"), cget(cn, d, "8월")
        st = wstat(mon, d)
        foot = []
        small = min(a.w25, a.w26, j.w25, j.w26) < 1.0
        if small: foot.append("†")
        if d == "평림": foot.append("ⓕ")
        if d == "남강": foot.append("ⓖ")
        if d == "대청": foot.append("ⓗ")
        dk = SMALL_DKM2.get(d, a.new - a.lost)
        dks = ("반올림 수준" if abs(dk) <= 0.02 else f2(dk, True)) if small else f2(a.w26 - a.w25, True)
        rows.append([d + (" ★" if d in DAMS8 else ""), f"{j.회차} ({f1(j.pct)})" if d == "평림" else f"{j.회차} {f1(j.pct)}",
                     f"{a.회차} {f1(a.pct)}", f"{a.w25:.2f} → {a.w26:.2f}", dks,
                     anom_s(st), where(st), anom_s(st, "a25"), " ".join(foot) or "—"])
        i = len(rows) - 1
        if d in DAMS8:
            cell[(i, 0)] = B; cell[(i, 2)] = B
        if a.pct > 0:
            cell[(i, 2)] = UP
        if d == "평림":
            cell[(i, 1)] = GREY
        if d == "대청":
            cell[(i, 2)] = GREY
        if st is not None and st["anom"] > 0.005:
            cell[(i, 5)] = UP; cell[(i, 6)] = UP
        if st is not None and not st["lower"]:
            cell[(i, 6)] = HI
    T.append(Tbl("표10_남부댐유역_전수_보조",
                 "Table 10. (보조) 남부 댐유역 전수 — SAR 수체 면적 Δ%(정본 쌍)와 그 댐의 WAMIS 10년 기록",
                 ["댐유역", "7월 쌍 · Δ%", "8월 쌍 · Δ%", "8월 수체 km²\n2025 → 2026", "Δkm²",
                  "WAMIS 2026-08\n중앙값 대비 (m)", "10년 기록에서\n2026-08", "2025-08\n중앙값 대비 (m)", "각주"],
                 rows,
                 notes=["★ 본문 8개. 쌍: 7월A ASC54 0725↔0715 · 7월B DESC134 0718↔0720 · 8월A ASC54 0806↔0802 · 8월B ASC127 0811↔0807. "
                        "두 회차에 모두 있는 유역은 ① 커버 최대 ② 부모 대권역과 같은 쌍 순으로 배정(gee canon_table 규칙).",
                        "WAMIS 열은 그 유역 댐을 자기 10년 기록과 비교한 값(Table 2 와 같다). 군위·보현산·보령은 WAMIS 월자료 없음.",
                        "† 수체 1 km² 미만(Δkm² 는 소실−신규, 0.02 이하는 반올림 수준이라 인용하지 않음; 보현산 −48.6% 는 −0.37 km²) · "
                        "ⓕ 평림 7월 임계 비일관(감소로 세지 않음) · ⓖ 남강: 커버 때문에 B군 배정, 임계는 낙동강(다른 궤도)에서 상속 · "
                        "ⓗ 대청: 궤도에 따라 −4.4 ↔ +0.5%, 저수지 −0.3% — 「변화 없음」으로 읽고 Δ% 를 인용하지 않는다."],
                 align=["l", "c", "c", "c", "c", "c", "l", "c", "c"], cell=cell))

    # 표11 ─ 금강
    res = pd.read_csv(GEE / "reservoir_km2_v11.csv", encoding="utf-8-sig").set_index("저수지")

    def rpct(name):
        a, b = res.loc[name, "2025-08-11_ASC127"], res.loc[name, "2026-08-07_ASC127"]
        return f"{f1(100 * (b / a - 1))} ({a:.2f}→{b:.2f})"

    rows, cell = [], {}
    for lab, kind, nm, rname, wd in (("용담", "댐유역", "용담", "용담호", "용담"), ("대청", "댐유역", "대청", "대청호", "대청"),
                                     ("보령", "댐유역", "보령", None, None), ("금 강 (대권역)", "대권역", "금 강", None, None),
                                     ("금강서해 (대권역)", "대권역", "금강서해", None, None),
                                     ("부안 (만경·동진)", "댐유역", "부안", None, "부안"),
                                     ("만경·동진 (대권역) ⓔ", "대권역", "만경동진", None, None)):
        j, a = cget(cn, nm, "7월", kind), cget(cn, nm, "8월", kind)
        st = wstat(mon, wd) if wd else None
        rows.append([lab, f1(j.pct), f1(a.pct), rpct(rname) if rname else "—",
                     (lvl(st) + " · " + where(st)) if st else "—",
                     {"용담": "WAMIS 10년 최저와 같은 방향 — 유역·저수지 1.7%p 안", "대청": "신호 없음 — Δ% 인용 안 함 ⓗ",
                      "부안": "약하지만 같은 방향", "만경동진": "하구·간척지 — 저수지와 다른 물"}.get(nm, "약한 감소")])
        if nm == "용담":
            cell[(len(rows) - 1, 2)] = HI; cell[(len(rows) - 1, 3)] = HI
        if nm == "대청":
            cell[(len(rows) - 1, None)] = {"color": "#7a7a7a"}
    T.append(Tbl("표11_금강권역_분해",
                 "Table 11. (보조) 금강 권역 — 공식 단계는 없지만 용담은 10년 최저·크게 감소, 대청은 변화 없음",
                 ["대상", "SAR 유역\n7월 Δ%", "SAR 유역\n8월 Δ%", "저수지 SAR 8월 Δ%\n(ASC127, km²)",
                  "WAMIS 2026-08 저수위\n(10년 중앙값 대비 m)", "해석"],
                 rows,
                 notes=["포스터 문구: 「금강은 8월 공식 단계 없음 · 용담은 WAMIS 10년 최저이고 SAR 로도 크게 감소, 대청은 변화 없음」. "
                        "대청은 궤도에 따라 −4.4 ↔ +0.5% 로 갈리고 8월B 두 해 임계가 한 구간 다르며(−20.5 ↔ −20.676) 만수 대비 면적이 "
                        "2025 94.1% · 2026 93.9% 로 둘 다 거의 만수다. 만경·동진은 금강과 별도 권역(대표 댐 부안)이고 부안호는 저수지 대조 목록에 없다."],
                 align=["l", "c", "c", "c", "c", "l"], cell=cell))
    return T


# ════════════════════════ 문서 ════════════════════════
FIGS = [
    ("그림1_WAMIS_8월저수위_10년추이.png", "① 이례성",
     "WAMIS 8월 평균 저수위(2017~2026) — 8개 저수지",
     "Fig. 1. 8개 저수지의 8월 평균 저수위(WAMIS, 2017~2026). 주황 = 2025→2026, 점선 = 그 댐의 10년 중앙값.",
     ["칸마다 세로축이 다르다(저수지마다 수위 높이가 100~360 m 로 다르다).",
      "2026년은 8곳 모두 자기 10년 중앙값 아래 — 안동·밀양·성덕·옥정호는 10년 최저, 임하·영천·운문은 2022년 다음으로 낮다.",
      "2025년(파랑)은 안동을 빼고 모두 중앙값 위 — 전년 대비 감소가 평년 대비보다 크게 보이는 이유."]),
    ("그림2_남부지도_WAMIS10년기록_SAR댐유역.png", "① + ②",
     "남부 지도 — (a) 댐별 10년 기록과 정부 단계 (b) SAR 댐유역 8월 Δ%",
     "Fig. 2. (a) 2026년 8월 WAMIS 저수위를 각 댐의 10년 기록과 비교한 결과와 8월 가뭄 예·경보 단계. "
     "(b) 같은 댐유역의 SAR 수체 면적 변화(2025→2026, 8월, 정본 1:1 쌍).",
     ["두 칸이 같은 유역 폴리곤이다 — 왼쪽은 공식 자료, 오른쪽은 SAR 관측.",
      "(a) 점 4개(동복·담양·장성·나주)는 유역 폴리곤이 없는 WAMIS 댐이다. 군위·보현산·보령은 WAMIS 월자료가 없다.",
      "영주·안계는 2026-08 저수위가 자기 중앙값 위이고 SAR 도 증가(+17.3·+10.8%)다 — 같은 방향."]),
    ("그림3_8개댐유역_SAR_지도막대.png", "②",
     "8개 댐유역 SAR 수체 면적 변화 — (a) 8월 지도 (b) 7·8월 막대",
     "Fig. 3. (a) 8개 댐유역의 8월 SAR 수체 면적 변화(2025→2026, 동일 상대궤도 1:1 쌍). "
     "(b) 같은 댐유역의 7·8월 Δ% — 괄호는 8월 가뭄 예·경보 단계.",
     ["현 원고 Fig. 1 과 같은 그림(안동 유역 북단이 잘리던 것을 고쳤다).", "평림 7월은 두 해 임계가 달라 그리지 않는다(ⓕ)."]),
    ("그림4_SAR면적_WAMIS저수량_민감도.png", "③",
     "SAR 면적 ↔ WAMIS 저수량 · 민감도 4쌍",
     "Fig. 4. (a) 저수지 SAR 면적 Δ% ↔ WAMIS 저수량 Δ%(음영 = 하한). (b) 8월 DESC134 S1C↔S1C 민감도 4쌍과 본쌍.",
     ["현 원고 Fig. 2 와 같은 그림. 32칸 중 30칸 하한, 민감도 20칸 모두 감소."]),
    ("그림5_SAR_S2_화소대응.png", "S2 교차검증",
     "SAR·S2 화소 대응 — 8월 공통 유효화소 85% 이상 4곳",
     "Fig. 5. 8월 85% 이상 4개 유역의 SAR·S2 수체 변화(본 저수지 100 m 띠, S2 는 MNDWI>0 고정 임계).",
     ["현 원고 Fig. 4 와 같은 그림. 일치율은 그림에 싣지 않고 본문에 영천·운문·밀양만 인용."]),
    ("그림6_관측쌍_타임라인_강우.png", "방법",
     "비교 설계 — 동일 상대궤도 1:1 관측쌍과 직전 5일 강우",
     "Fig. 6. 세 상대궤도 × 두 해의 Sentinel-1 관측과 비교에 쓴 1:1 쌍. 숫자는 쌍 대상 권역의 직전 5일 강우(권역 중앙값, mm).",
     ["2026 쪽은 거의 건조하고 2025 쪽에 비 온 직후가 많다(0718 · 0806 · 0811).", "강우는 쌍을 거르는 기준이 아니라 해석용 주석이다."]),
    ("그림7_부록_강우위상_저수지vs유역.png", "④ 부록",
     "강우 위상 — 저수지는 둔하고 유역(하천)은 민감하다",
     "부록 그림. 2025 날짜를 비 온 0806 → 마른 0812 로 바꿨을 때의 SAR Δ% 변화(2026 = 0802 고정).",
     ["저수지 ≤1.0%p, 유역은 안동·임하에서 4.9·6.2%p — 차이는 담수호 밖 하천에서 난다.", "밀양·운문은 0812 도 젖어 건조 대조가 아니다(회색)."]),
]

DRAFT_FIX = [
    ("Introduction 글머리", "「WAMIS 저수위의 10년 순위로 남부 대권역(…)의 2026년 여름 가뭄 이례성을 확인」",
     "「WAMIS 댐별 10년 저수위 기록과 정부 가뭄 예·경보로 2026년 여름 남부 가뭄의 이례성을 확인」"),
    ("Table 1 (원고)", "10년 중앙값·순위 열 — 「10년 중앙값」 은 10개 중 6번째 값이라 표준 중앙값과 0.1~1.3 m 다르다",
     "그림1(꺾은선) 또는 Table 2 로 바꾼다. 순위 열은 빼고 「10년 최저 / ○○년 다음으로 낮음」 으로"),
    ("KPI 「7/8」", "「8월 저수위 10년 최저·2위 (WAMIS · Table 1)」",
     "「8/8 — 2026년 8월 저수위가 자기 10년 중앙값 아래 (4곳 10년 최저)」"),
    ("결과 표 아래 각주", "「2025년 8월은 8곳 중 6곳이 평년 위라 …」",
     "「2025년 8월은 안동을 뺀 7곳이 자기 10년 중앙값보다 높아 …」 (6곳은 오기)"),
    ("결과 표 아래 각주 ⓒ", "「안동은 2025년도 10년 중 2위로 낮아 …」", "「안동은 2025년도 10년 중앙값보다 2.4 m 낮아 …」"),
    ("Results 글머리(민감도)", "「위성 차 0.5%, 임계 ±1구간 ≤1.1%p」",
     "「… 임계 ±1구간은 맑은 4유역에서 ≤1.1%p」 — 임하는 1.30%p(상대 18%)라 8개 전체에 쓰지 않는다"),
    ("Conclusion 첫 줄", "「WAMIS 10년 저수위 순위와 정부 가뭄 예·경보로 확인되는 …」",
     "「WAMIS 댐별 10년 저수위 기록과 정부 가뭄 예·경보로 확인되는 …」"),
]

DONT = ["순위·권역 순위 중앙값으로 가뭄 권역을 판정하는 것 — WAMIS 는 댐별 10년 기록, 정부는 보도자료 자체로",
        "「10년 중 N위」 표현 → 「10년 최저」 · 「○○년 다음으로 낮음」 · 「중앙값보다 ○ m 낮음」",
        "v10 이하 수치(EGM 끔 판)", "map_area_s2.csv 광학 수면적", "_diag_* 진단 파일", "접경·한강 권역(논문으로)",
        "「모두 하한」 — 32칸 중 30칸", "S2 「파장 순서」(STAC 16사례 0/16)", "대청 Δ%", "성덕 화소 대응 30.2%",
        "reservoir_km2 의 ·궤도혼합 열", "수체 1 km² 미만의 반올림 수준 Δkm²(선암·감포·대암 등)",
        "「4유역 독립 확인」 — 같은 영상·같은 임계", "유역평균 ΔMNDWI 로 수면 말하기", "평림 7월 −1.5% 와 「두 달 다 8/8」"]


def write_doc(out: Path, tbls: list) -> Path:
    L = ["# 2026 남부 가뭄 포스터 — S1 SAR 그림·표 정리", "",
         "2026 원격탐사학회 추계학술대회 포스터(KSRS-131)용. 자료는 SAR v11(EGM 켬·POEORB), 범위는 남부(낙동강·금강·섬진강·영산강 계열 "
         "15개 대권역). 근거: gee `docs/DROUGHT_2026_EVIDENCE_KR.md` 「포스터 조립용 최종 목록」.", "",
         "## 0. 판단 기준 — 자료마다 자기 기준으로", "",
         "| 자료 | 무엇으로 판단하나 | 표현 |", "|:--|:--|:--|",
         "| **WAMIS 저수위** | 댐마다 **자기 10년(2017~2026)** 월평균 저수위와 비교한다. 다른 댐·권역과 순위를 매겨 가뭄 권역을 정하지 않는다. "
         "| 2026 값, 10년 중앙값 대비(m), 10년 최저 / ○○년 다음으로 낮음 |",
         "| **정부 보도자료** | **보도자료 자체의 판단**을 그대로 쓴다. 그 판단도 평년·예년 대비다(댐 저수량 = 예년, 강수·농업용 저수지 = 평년). "
         "| 행안부 「8월 가뭄 예·경보」 단계(8.10 발표), 기후부 단계 진입일 |",
         "| **SAR** | 판정이 아니라 **관측** — 같은 상대궤도 1:1 쌍의 전년도(2025) 대비 수체 면적 변화, 저수량 감소의 하한 | Δ%, Δkm² |",
         "", "두 공식 자료가 다르게 말하면(예: 용담 — WAMIS 10년 최저인데 공식 단계 없음) 각자 기준으로 따로 적는다.", "",
         "## 1. 포스터 서사와 배치", "",
         "| 서사 | 그림 | 표 |", "|:--|:--|:--|",
         "| ① 2026 여름 남부 가뭄은 이례적이다 | 그림1 · 그림2(a) | Table 1 · 2 · 4 |",
         "| ② SAR 이 그 감소를 보여준다 | 그림2(b) · 그림3 | Table 1 · 3 · 10 |",
         "| ③ 공식 자료와 맞는다(하한) | 그림4 | Table 5 · 6 |",
         "| S2 교차검증 | 그림5 (+ S2_자료 폴더 그림1·3) | Table 7 |",
         "| 방법 | 그림6 | Table 8 |",
         "| ④ 부록 | 그림7 | Table 9 · 11 |", "",
         "## 2. 그림", ""]
    for i, (fn, role, title, cap, bullets) in enumerate(FIGS, 1):
        L += [f"### 그림{i}. {title}  `{role}`", "", f"![그림{i}]({fn})", "", f"**캡션 초안** — {cap}", ""]
        L += [f"- {b}" for b in bullets] + [""]
    L += ["## 3. 표", "",
          "같은 표가 PNG(`표N_….png`, 보기용) · `표모음_편집용.xlsx` · `표모음_A0_편집용.pptx`(원고와 같은 A0 세로, 머리 22 pt · 본문 20 pt, "
          "복사해 붙이면 편집 가능)로 있다. 아래 표는 그 파일들과 같은 값에서 만들었다.", ""]
    for t in tbls:
        L += [f"### {t.caption}", "", md_table(t), ""] + [f"- {n}" for n in t.notes] + [""]
    L += ["## 4. 현재 원고(`KSRS-131_유하은_v11수정_Results.pptx`)에서 고칠 것", "",
          "| 위치 | 지금 | 고칠 문구 |", "|:--|:--|:--|"]
    L += [f"| {a} | {b} | {c} |" for a, b, c in DRAFT_FIX] + [""]
    L += ["## 5. 포스터에 쓰지 말 것", ""] + [f"- {d}" for d in DONT] + [""]
    L += ["## 6. 파일·재현·출처", "",
          "- 이 폴더의 `poster_A/B/C_*.png`·`poster_S2_fig1_crop.png`·CSV 는 이전 세션 산출이라 그대로 두었다(그림3·4·5 와 같은 그림).",
          "- 재현: `$env:GDAL_DATA = \"F:\\envs\\gis_clean\\Library\\share\\gdal\"` 후 "
          "`F:\\envs\\gis_clean\\python.exe F:\\06_SAR_system\\S1\\scripts\\poster_kit.py --out <이 폴더> --pylib <openpyxl·python-pptx 경로>` "
          "(스크립트 사본 `_재현\\`).",
          f"- WAMIS 월자료(남부 33개 댐 7·8월, 2017~2026): `_자료\\{MON_NAME}` — 지우면 다시 받는다.",
          "- 자동 대조: 유역 Δ% 88칸 = gee `canon_table.txt`.",
          "- 유역 Δ%: gee `Korea_WaterDetection_2025_2026/output/watershed_change_v11_{jul_a,jul_b,aug_a,aug_b}[_fix].csv` · "
          "저수지 면적: gee `reservoir_km2_v11*.csv`(날짜 열) · SAR↔WAMIS: S1 `experiments/swot_wamis/sar_vs_wamis_v11{,_sens}.csv` · "
          "강우: S1 `data/rain_phase_5d.csv` · S2: `..\\S2_자료_20261015\\표1·표2` · 공식: `..\\행정안전부 보도자료\\` · `..\\기후에너지환경부 보도자료\\`.",
          ""]
    p = out / "00_S1_SAR_포스터_그림표_정리.md"
    p.write_text("\n".join(L), encoding="utf-8")
    return p


# ════════════════════════ 실행 ════════════════════════
def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--out", type=Path, required=True)
    ap.add_argument("--pylib", type=Path, help="openpyxl·python-pptx 가 있는 site-packages (뒤에 붙인다)")
    ap.add_argument("--only", nargs="*", help="그림1 … 그림7 · 표 (생략하면 전부 + 문서)")
    args = ap.parse_args()
    if args.pylib:
        sys.path.append(str(args.pylib))
    psf.style()
    out = args.out
    out.mkdir(parents=True, exist_ok=True)
    todo = set(args.only or ["그림1", "그림2", "그림3", "그림4", "그림5", "그림6", "그림7", "표"])

    cn = canon()
    mon = wamis_south(out / "_자료" / MON_NAME)
    sw = pd.read_csv(WAMIS_DIR / "sar_vs_wamis_v11.csv", encoding="utf-8-sig")
    swsens = pd.read_csv(WAMIS_DIR / "sar_vs_wamis_v11_sens.csv", encoding="utf-8-sig")
    sens = psf.sens_table()
    rain = pd.read_csv(RAIN_CSV, dtype={"date": str, "orbit": str})

    if "그림1" in todo:
        fig1(mon, out / FIGS[0][0])
    if "그림2" in todo:
        fig2(cn, mon, out / FIGS[1][0])
    if "그림3" in todo:
        psf.fig_a(psf.basin_table(), out / FIGS[2][0])
    if "그림4" in todo:
        main_w, sens_w = psf.wamis_tables()
        psf.fig_b(main_w, sens_w, sens, out / FIGS[3][0])
    if "그림5" in todo:
        psf.fig_c(out / FIGS[4][0])
    if "그림6" in todo:
        fig6(rain, out / FIGS[5][0])
    if "그림7" in todo:
        fig7(cn, out / FIGS[6][0])
    if "표" in todo:
        tbls = build_tables(cn, mon, sw, swsens, sens)
        widths = {t.key: render_png(t, out / f"{t.key}.png") for t in tbls}
        try:
            write_xlsx(tbls, out / "표모음_편집용.xlsx")
            write_pptx(tbls, widths, out / "표모음_A0_편집용.pptx")
        except ImportError as e:
            print(f"⚠ xlsx·pptx 생략 — {e} (--pylib 로 openpyxl·python-pptx 경로를 주면 만든다)")
        print(f"표 {len(tbls)}개 → 문서 {write_doc(out, tbls).name}")
    print(f"→ {out}")


if __name__ == "__main__":
    main()
